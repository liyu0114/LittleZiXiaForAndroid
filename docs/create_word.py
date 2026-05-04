#!/usr/bin/env python3
# -*- coding: utf-8 -*-

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("需要安装 python-docx: pip install python-docx")
    exit(1)

doc = Document()

# 标题
title = doc.add_heading('OpenClaw 移动端客户端 - 项目报告', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 基本信息
doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('版本: ').bold = True
p.add_run('v1.5.1')
p = doc.add_paragraph()
p.add_run('构建日期: ').bold = True
p.add_run('2026-03-20')
p = doc.add_paragraph()
p.add_run('APK 文件: ').bold = True
p.add_run('OpenClaw_v1.5.1_20260320_0750.apk')

doc.add_paragraph()
doc.add_heading('一、项目概述', level=1)

doc.add_paragraph(
    'OpenClaw 移动端客户端是一个 Flutter 应用，用于远程连接和控制 OpenClaw Gateway。'
    '用户可以通过手机与 Gateway 上的 AI Agent 进行双向通信，发送指令并接收响应。'
)

doc.add_heading('核心功能', level=2)
features = [
    'WebSocket 双向通信 - 实时消息收发',
    'HTTP API 支持 - 消息发送和状态查询',
    '任务状态追踪 - 实时显示 AI 执行步骤',
    '配置管理 - Gateway URL 和 Token 持久化存储',
    'Material Design 3 - 现代化 UI 设计',
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('二、技术架构', level=1)

doc.add_heading('技术栈', level=2)
table = doc.add_table(rows=6, cols=2)
table.style = 'Table Grid'
data = [
    ('组件', '技术'),
    ('框架', 'Flutter 3.x'),
    ('语言', 'Dart'),
    ('状态管理', 'Provider'),
    ('本地存储', 'SharedPreferences'),
    ('网络通信', 'HTTP + WebSocket'),
]
for i, (a, b) in enumerate(data):
    table.rows[i].cells[0].text = a
    table.rows[i].cells[1].text = b
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_heading('三、安装与使用', level=1)

doc.add_heading('系统要求', level=2)
doc.add_paragraph('Android 5.0 (API 21) 或更高版本')
doc.add_paragraph('网络连接（局域网或互联网）')

doc.add_heading('安装步骤', level=2)
steps = [
    '传输 APK 到 Android 设备',
    '允许安装未知来源应用',
    '安装 APK',
    '打开应用',
]
for i, s in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {s}')

doc.add_heading('配置步骤', level=2)
steps = [
    '在"配置"标签页输入 Gateway URL',
    'Token 已内置，无需修改',
    '点击"连接"按钮',
    '切换到"会话"标签页开始对话',
]
for i, s in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {s}')

doc.add_heading('四、已知问题', level=1)

doc.add_heading('Gateway 消息推送', level=2)
p = doc.add_paragraph()
p.add_run('现象: ').bold = True
p.add_run('APP 发送消息成功，但可能收不到 Agent 的回复')
p = doc.add_paragraph()
p.add_run('原因: ').bold = True
p.add_run('Gateway 需要实现 WebSocket 客户端消息路由机制')
p = doc.add_paragraph()
p.add_run('临时方案: ').bold = True
p.add_run('使用 HTTP 轮询或等待 Gateway 更新')

doc.add_heading('五、后续计划', level=1)
plans = [
    '自动重连机制',
    '消息推送优化',
    '多 Gateway 支持',
    '消息历史同步',
    '深色模式切换',
]
for p in plans:
    doc.add_paragraph(p, style='List Bullet')

doc.add_paragraph()
doc.add_paragraph('—' * 30)
p = doc.add_paragraph()
p.add_run('文档版本: ').bold = True
p.add_run('1.0')
p = doc.add_paragraph()
p.add_run('最后更新: ').bold = True
p.add_run('2026-03-20')

# 保存
doc.save(r'D:\mobile-client\docs\OpenClaw_Mobile_Report.docx')
print('Word 文档已生成: D:\\mobile-client\\docs\\OpenClaw_Mobile_Report.docx')
